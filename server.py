#!/usr/bin/env python3
import json, os, urllib.request, urllib.error, urllib.parse, subprocess, platform, ssl, io, re, base64
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path

ROOT = Path(__file__).resolve().parent
HOST = '127.0.0.1'
PORT = int(os.environ.get('PDF_TRANSLATOR_PORT', '8765'))
APP_ID = 'scholar-pdf-translator'
APP_VERSION = '2.6.0'
PROVIDERS = {
    'deepseek': {'label':'DeepSeek','env':'DEEPSEEK_API_KEY','base':'https://api.deepseek.com','model':'deepseek-chat','kind':'openai'},
    'openai': {'label':'OpenAI','env':'OPENAI_API_KEY','base':'https://api.openai.com/v1','model':'gpt-5','kind':'openai'},
    'anthropic': {'label':'Anthropic Claude','env':'ANTHROPIC_API_KEY','base':'https://api.anthropic.com','model':'claude-sonnet-4-5','kind':'anthropic'},
    'gemini': {'label':'Google Gemini','env':'GEMINI_API_KEY','base':'https://generativelanguage.googleapis.com/v1beta','model':'gemini-2.5-pro','kind':'gemini'},
    'qwen': {'label':'Alibaba Qwen','env':'DASHSCOPE_API_KEY','base':'https://dashscope.aliyuncs.com/compatible-mode/v1','model':'qwen-plus','kind':'openai'},
    'kimi': {'label':'Moonshot Kimi','env':'MOONSHOT_API_KEY','base':'https://api.moonshot.cn/v1','model':'moonshot-v1-32k','kind':'openai'},
    'openrouter': {'label':'OpenRouter','env':'OPENROUTER_API_KEY','base':'https://openrouter.ai/api/v1','model':'openai/gpt-5','kind':'openai'},
    'custom': {'label':'OpenAI-compatible custom','env':'PDF_TRANSLATOR_API_KEY','base':'','model':'','kind':'openai'},
}
KEYCHAIN_ACCOUNT = os.environ.get('USER') or os.environ.get('LOGNAME') or os.environ.get('USERNAME') or 'local-user'
KEY_DIR = Path.home() / '.scholar_pdf_translator'


def build_ssl_context():
    """Build a verified TLS context, preferring certifi's Mozilla CA bundle."""
    try:
        import certifi
        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        # Fall back to the Python/system trust store. We deliberately do not
        # disable certificate verification.
        return ssl.create_default_context()

def _windows_powershell():
    for name in ('powershell.exe', 'powershell'):
        try:
            r = subprocess.run([name, '-NoProfile', '-Command', '$PSVersionTable.PSVersion.Major'], capture_output=True, text=True, timeout=5)
            if r.returncode == 0:
                return name
        except Exception:
            pass
    return 'powershell.exe'


def normalize_provider(provider):
    p=(provider or 'deepseek').strip().lower()
    return p if p in PROVIDERS else 'custom'


def provider_info(provider):
    return PROVIDERS[normalize_provider(provider)]


def _safe_provider_name(provider):
    return re.sub(r'[^a-z0-9_-]+','_',normalize_provider(provider))


def _keychain_service(provider):
    return 'Scholar PDF Translator - ' + _safe_provider_name(provider)


def _fallback_key_file(provider):
    return KEY_DIR / (_safe_provider_name(provider) + '_api_key')


def _windows_key_file(provider):
    return KEY_DIR / (_safe_provider_name(provider) + '_api_key.dpapi')


def load_saved_api_key(provider='deepseek'):
    provider=normalize_provider(provider); system=platform.system()
    if system == 'Darwin':
        try:
            r = subprocess.run(['security','find-generic-password','-s',_keychain_service(provider),'-a',KEYCHAIN_ACCOUNT,'-w'], capture_output=True, text=True, timeout=10)
            if r.returncode == 0: return r.stdout.strip()
        except Exception: pass
        return ''
    if system == 'Windows':
        target=_windows_key_file(provider)
        if not target.exists(): return ''
        ps=_windows_powershell()
        script=r'''$p=$env:PDF_TRANSLATOR_KEY_FILE; if (!(Test-Path -LiteralPath $p)) { exit 3 }; $enc=Get-Content -LiteralPath $p -Raw; $sec=ConvertTo-SecureString $enc; $b=[Runtime.InteropServices.Marshal]::SecureStringToBSTR($sec); try {[Runtime.InteropServices.Marshal]::PtrToStringBSTR($b)} finally {[Runtime.InteropServices.Marshal]::ZeroFreeBSTR($b)}'''
        env=os.environ.copy(); env['PDF_TRANSLATOR_KEY_FILE']=str(target)
        try:
            r=subprocess.run([ps,'-NoProfile','-NonInteractive','-Command',script],capture_output=True,text=True,timeout=15,env=env)
            if r.returncode==0:return r.stdout.strip()
        except Exception: pass
        return ''
    try:return _fallback_key_file(provider).read_text(encoding='utf-8').strip()
    except Exception:return ''


def save_api_key(api_key, provider='deepseek'):
    api_key=(api_key or '').strip(); provider=normalize_provider(provider)
    if not api_key: raise ValueError('API Key 不能为空')
    system=platform.system()
    if system=='Darwin':
        r=subprocess.run(['security','add-generic-password','-U','-s',_keychain_service(provider),'-a',KEYCHAIN_ACCOUNT,'-w',api_key],capture_output=True,text=True,timeout=10)
        if r.returncode!=0: raise RuntimeError((r.stderr or r.stdout or '无法写入钥匙串').strip())
        return 'macOS Keychain'
    if system=='Windows':
        target=_windows_key_file(provider); target.parent.mkdir(parents=True,exist_ok=True); ps=_windows_powershell()
        script=r'''$sec=ConvertTo-SecureString -String $env:PDF_TRANSLATOR_API_KEY -AsPlainText -Force; $enc=$sec | ConvertFrom-SecureString; Set-Content -LiteralPath $env:PDF_TRANSLATOR_KEY_FILE -Value $enc -Encoding UTF8'''
        env=os.environ.copy(); env['PDF_TRANSLATOR_API_KEY']=api_key; env['PDF_TRANSLATOR_KEY_FILE']=str(target)
        r=subprocess.run([ps,'-NoProfile','-NonInteractive','-Command',script],capture_output=True,text=True,timeout=15,env=env)
        if r.returncode!=0: raise RuntimeError((r.stderr or r.stdout or '无法写入 Windows DPAPI 凭据').strip())
        return 'Windows DPAPI'
    target=_fallback_key_file(provider); target.parent.mkdir(parents=True,exist_ok=True); target.write_text(api_key,encoding='utf-8')
    try:os.chmod(target,0o600)
    except Exception:pass
    return 'local secure file'


def delete_saved_api_key(provider='deepseek'):
    provider=normalize_provider(provider); system=platform.system()
    if system=='Darwin':
        r=subprocess.run(['security','delete-generic-password','-s',_keychain_service(provider),'-a',KEYCHAIN_ACCOUNT],capture_output=True,text=True,timeout=10)
        if r.returncode not in (0,44): raise RuntimeError((r.stderr or r.stdout or '无法删除钥匙串项目').strip())
        return
    target=_windows_key_file(provider) if system=='Windows' else _fallback_key_file(provider)
    try:target.unlink(missing_ok=True)
    except TypeError:
        if target.exists():target.unlink()


def _system_curl():
    if platform.system()=='Windows': return 'curl.exe'
    return '/usr/bin/curl' if Path('/usr/bin/curl').exists() else 'curl'


def call_json_via_curl(url, headers, payload, provider_label='AI'):
    body=json.dumps(payload,ensure_ascii=False)
    cmd=[_system_curl(),'--silent','--show-error','--connect-timeout','20','--max-time','300','--request','POST','--header','Content-Type: application/json']
    for k,v in headers.items(): cmd += ['--header', f'{k}: {v}']
    cmd += ['--data-binary','@-','--write-out','\n__HTTP_STATUS__:%{http_code}',url]
    r=subprocess.run(cmd,input=body,capture_output=True,text=True,timeout=310)
    if r.returncode!=0:
        msg=(r.stderr or r.stdout or '').strip(); raise RuntimeError(f'系统 curl 连接 {provider_label} 失败：'+(msg[:1600] or f'退出码 {r.returncode}'))
    marker='\n__HTTP_STATUS__:'
    if marker not in r.stdout: raise RuntimeError(f'{provider_label} 返回格式异常：未取得 HTTP 状态码')
    response_text,status_text=r.stdout.rsplit(marker,1)
    try:status=int(status_text.strip())
    except ValueError: raise RuntimeError(f'{provider_label} 返回格式异常：HTTP 状态码无效')
    if not (200<=status<300): raise RuntimeError(f'{provider_label} HTTP {status}: {response_text[:1800]}')
    try:return json.loads(response_text)
    except json.JSONDecodeError as e: raise RuntimeError(f'{provider_label} 返回了无法解析的 JSON：'+response_text[:1200]) from e


def call_json_via_python(url, headers, payload, provider_label='AI'):
    req=urllib.request.Request(url,data=json.dumps(payload,ensure_ascii=False).encode('utf-8'),headers={'Content-Type':'application/json',**headers},method='POST')
    try:
        with urllib.request.urlopen(req,timeout=300,context=build_ssl_context()) as r:return json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        detail=e.read().decode(errors='replace'); raise RuntimeError(f'{provider_label} HTTP {e.code}: {detail[:1500]}')


def _post_json(url, headers, payload, provider_label):
    if platform.system() in ('Darwin','Windows'):
        return call_json_via_curl(url,headers,payload,provider_label)
    try:return call_json_via_curl(url,headers,payload,provider_label)
    except Exception:return call_json_via_python(url,headers,payload,provider_label)


def _join_url(base, suffix):
    base=(base or '').rstrip('/')
    if base.endswith(suffix): return base
    return base+suffix


def _user_prompt(data):
    text=data.get('text') or ''; context=(data.get('context') or '').strip(); source=(data.get('source_text') or '').strip(); task=data.get('task','translate')
    if task=='proofread':
        user=f'SOURCE:\n{source}\n\nTRANSLATION TO PROOFREAD:\n{text}'
        if context:user=f'上一页译文末尾（仅供术语连续性参考）：\n{context}\n\n'+user
    elif task=='reconstruct':
        user=f'待重建的中文译文：\n{text}'
        if source:user+=f'\n\n原文对照（仅用于判断标题层级、注号、注释边界和 OCR 污染；不要重新翻译）：\n{source}'
        if context:user=f'前文结构上下文（只用于章节连续性）：\n{context}\n\n'+user
    else:
        user=text
        if context:user=f'前文上下文（只用于术语和衔接，不要重复输出）：\n{context}\n\n待处理文本：\n{text}'
    return user


def call_ai(data):
    provider=normalize_provider(data.get('provider')); info=provider_info(provider); label=info['label']
    env_key=os.environ.get(info.get('env',''),'') if info.get('env') else ''
    api_key=(data.get('api_key') or env_key or load_saved_api_key(provider) or '').strip()
    if not api_key: raise ValueError(f'缺少 {label} API Key')
    base=(data.get('api_base') or info.get('base') or '').rstrip('/')
    model=(data.get('model') or info.get('model') or '').strip()
    if not base: raise ValueError(f'{label} API Base 不能为空')
    if not model: raise ValueError(f'{label} 模型名称不能为空')
    system=system_prompt(data); user=_user_prompt(data); kind=info.get('kind','openai')
    if kind=='anthropic':
        url=_join_url(base,'/v1/messages')
        payload={'model':model,'max_tokens':8192,'temperature':0.15,'system':system,'messages':[{'role':'user','content':user}]}
        obj=_post_json(url,{'x-api-key':api_key,'anthropic-version':'2023-06-01'},payload,label)
        parts=obj.get('content') or []; text=''.join((p.get('text') or '') for p in parts if isinstance(p,dict) and p.get('type')=='text').strip()
        if not text: raise RuntimeError(f'{label} 返回中没有可用文本')
        return text
    if kind=='gemini':
        url=_join_url(base,f'/models/{model}:generateContent')
        payload={'systemInstruction':{'parts':[{'text':system}]},'contents':[{'role':'user','parts':[{'text':user}]}],'generationConfig':{'temperature':0.15}}
        obj=_post_json(url,{'x-goog-api-key':api_key},payload,label)
        candidates=obj.get('candidates') or []
        if not candidates: raise RuntimeError(f'{label} 返回中没有 candidates：{json.dumps(obj,ensure_ascii=False)[:900]}')
        parts=((candidates[0].get('content') or {}).get('parts') or []); text=''.join((p.get('text') or '') for p in parts if isinstance(p,dict)).strip()
        if not text: raise RuntimeError(f'{label} 返回中没有可用文本')
        return text
    url=_join_url(base,'/chat/completions')
    payload={'model':model,'messages':[{'role':'system','content':system},{'role':'user','content':user}],'temperature':0.15,'stream':False}
    obj=_post_json(url,{'Authorization':'Bearer '+api_key},payload,label)
    try:return (obj['choices'][0]['message']['content'] or '').strip()
    except Exception: raise RuntimeError(f'{label} 返回中没有 choices[0].message.content：{json.dumps(obj,ensure_ascii=False)[:900]}')


def _paras(text):
    text=(text or '').replace('\r\n','\n').replace('\r','\n').strip()
    if not text:
        return []
    # Translation prompts preserve paragraphs with blank lines. If a page contains
    # only line breaks, keep those lines together rather than exploding every line.
    parts=[re.sub(r'\s*\n\s*', ' ', x).strip() for x in re.split(r'\n\s*\n+', text) if x.strip()]
    return parts or [text]


def _looks_heading(text):
    t=re.sub(r'\s+',' ',(text or '').strip())
    if not t or len(t)>120:
        return False
    if re.match(r'^(chapter|part|section|introduction|conclusion|preface|foreword|contents|bibliography|references|appendix)\b', t, re.I):
        return True
    if re.match(r'^\d+(?:\.\d+){0,3}[\s、.)]+\S+', t):
        return True
    if re.match(r'^[一二三四五六七八九十百]+[、.．]\s*\S+', t):
        return True
    letters=[c for c in t if c.isalpha()]
    if letters and len(t)<80 and sum(c.isupper() for c in letters)/len(letters)>.82:
        return True
    # Common book/article heading shape: short, single-sentence line without end punctuation.
    if len(t)<=45 and not re.search(r'[,，。！？.!?;；:：]', t) and len(t.split())<=9:
        return True
    return False


def _join_across_page(prev, nxt, target=False):
    if not prev or not nxt or _looks_heading(prev) or _looks_heading(nxt):
        return False
    p=prev.rstrip(); n=nxt.lstrip()
    if target:
        # Chinese prose ending with a sentence terminator normally closes the paragraph.
        # Comma/semicolon/colon or no punctuation at a PDF page boundary is usually continuation.
        return not bool(re.search(r'[。！？!?][”’\"\']?$', p))
    # Strong signal for English/European prose: previous page lacks terminal punctuation,
    # or the next page begins with a lowercase continuation.
    if re.search(r'[.!?][”’\"\']?$', p):
        return False
    if n and n[0].islower():
        return True
    return not bool(re.search(r'[:;][”’\"\']?$', p))


def reflow_blocks(pages, field='target', target=True):
    blocks=[]
    for page in pages or []:
        ps=_paras(page.get(field) or (page.get('source') if field=='target' else ''))
        if not ps:
            continue
        start=len(blocks)
        for x in ps:
            kind='heading' if _looks_heading(x) else 'body'
            blocks.append({'text':x,'kind':kind,'page':page.get('n')})
        # Heal only the boundary created by the original PDF page break.
        if start>0 and start<len(blocks):
            a=blocks[start-1]; b=blocks[start]
            if a['kind']=='body' and b['kind']=='body' and _join_across_page(a['text'],b['text'],target=target):
                sep='' if target and re.search(r'[\u3400-\u9fff]$',a['text']) else ' '
                a['text']=a['text'].rstrip()+sep+b['text'].lstrip()
                del blocks[start]
    return blocks


def _set_run_font(run, latin='Times New Roman', east_asia='宋体', size_pt=None, bold=None):
    from docx.oxml.ns import qn
    if size_pt is not None:
        from docx.shared import Pt
        run.font.size=Pt(size_pt)
    run.font.name=latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia)
    if bold is not None:
        run.bold=bold
    try:
        from docx.shared import RGBColor
        run.font.color.rgb=RGBColor(0,0,0)
    except Exception:
        pass


def _style_doc(doc, bilingual=False):
    from docx.shared import Cm, Pt
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    sec=doc.sections[0]
    if bilingual:
        sec.orientation=WD_ORIENT.LANDSCAPE
        sec.page_width,sec.page_height=sec.page_height,sec.page_width
        sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.6); sec.left_margin=Cm(1.5); sec.right_margin=Cm(1.5)
    else:
        sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.4); sec.right_margin=Cm(2.4)
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); normal.font.size=Pt(11)
    pf=normal.paragraph_format; pf.line_spacing=1.5; pf.space_after=Pt(4)
    for name, size in [('Title',18),('Heading 1',15),('Heading 2',13)]:
        st=doc.styles[name]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); st.font.size=Pt(size)
    return sec


def build_translation_docx(data, bilingual=False):
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError('缺少 python-docx。请重新运行“启动翻译工具.command”，新版会自动安装该组件。') from e

    pages=data.get('pages') or []
    if not pages:
        raise ValueError('没有可导出的页面')
    title=(data.get('title') or data.get('file_name') or '翻译稿').strip()
    title=re.sub(r'\.pdf$','',title,flags=re.I)
    keep_page_markers=bool(data.get('keep_page_markers', False))
    doc=Document(); _style_doc(doc,bilingual=bilingual)
    cp=doc.core_properties; cp.title=title; cp.subject='Scholar PDF Translator 导出稿'

    tp=doc.add_paragraph(); tp.style='Title'; tp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    tr=tp.add_run(title); _set_run_font(tr,east_asia='黑体',size_pt=18,bold=True)

    if not bilingual:
        blocks=reflow_blocks(pages,'target',target=True)
        last_page=None
        for b in blocks:
            if keep_page_markers and b.get('page')!=last_page:
                m=doc.add_paragraph(); m.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                r=m.add_run(f'原 PDF 第 {b.get("page")} 页'); _set_run_font(r,east_asia='宋体',size_pt=8); r.font.italic=True
                last_page=b.get('page')
            if b['kind']=='heading':
                p=doc.add_paragraph(style='Heading 1')
                p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
                r=p.add_run(b['text']); _set_run_font(r,east_asia='黑体',size_pt=14,bold=True)
            else:
                p=doc.add_paragraph()
                p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent=Pt(22)
                p.paragraph_format.line_spacing=1.5
                r=p.add_run(b['text']); _set_run_font(r,east_asia='宋体',size_pt=11)
    else:
        note=doc.add_paragraph(); note.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=note.add_run('原文 / 译文对照版'); _set_run_font(rr,east_asia='黑体',size_pt=10,bold=True)
        table=doc.add_table(rows=1, cols=2)
        table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        widths=[Cm(12.6),Cm(12.6)]
        hdr=table.rows[0].cells
        for i,txt in enumerate(('Original','中文译文')):
            hdr[i].width=widths[i]; hdr[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=hdr[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(txt); _set_run_font(r,east_asia='黑体',size_pt=10,bold=True)
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'EDEDED'); hdr[i]._tc.get_or_add_tcPr().append(shd)
        for page in pages:
            src=_paras(page.get('source','')); tgt=_paras(page.get('target') or page.get('source',''))
            if keep_page_markers:
                cells=table.add_row().cells; merged=cells[0].merge(cells[1]); merged.text=''
                p=merged.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(f'原 PDF 第 {page.get("n")} 页'); _set_run_font(r,east_asia='宋体',size_pt=8); r.font.italic=True
            count=max(len(src),len(tgt),1)
            for i in range(count):
                cells=table.add_row().cells
                for j,text in enumerate((src[i] if i<len(src) else '', tgt[i] if i<len(tgt) else '')):
                    cells[j].width=widths[j]; cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                    p=cells[j].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.25
                    r=p.add_run(text); _set_run_font(r,east_asia='宋体',size_pt=9.5)
        # Repeat header row on subsequent pages.
        trPr=table.rows[0]._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

    # Footer page numbers using a PAGE field.
    for sec in doc.sections:
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)
    out=io.BytesIO(); doc.save(out); return out.getvalue()


NOISE_PATTERNS = [
    r'^\s*\d{1,5}\s*$',
    r'\b(barcode|call\s*number|library\s*copy|digitized\s*by|scanned\s*by|google\s*books|hathitrust|internet\s*archive)\b',
    r'(图书馆.{0,8}(藏|馆藏|借阅|索书)|馆藏章|扫描.{0,6}(制作|来源|编号)|条形码|条码|索书号)',
    r'^[\W_]{3,}$',
]


def _hard_noise(text):
    t=re.sub(r'\s+',' ',(text or '').strip())
    if not t:
        return True
    if len(t) <= 3 and not re.search(r'[\u3400-\u9fffA-Za-z0-9]', t):
        return True
    return any(re.search(p,t,re.I) for p in NOISE_PATTERNS)


def _set_update_fields(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    settings=doc.settings._element
    node=settings.find(qn('w:updateFields'))
    if node is None:
        node=OxmlElement('w:updateFields'); settings.append(node)
    node.set(qn('w:val'),'true')


def _add_toc_field(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p=doc.add_paragraph()
    p.alignment=1
    r=p.add_run('目录')
    _set_run_font(r,east_asia='黑体',size_pt=16,bold=True)
    p.paragraph_format.space_after = __import__('docx').shared.Pt(10)
    toc=doc.add_paragraph()
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' TOC \\o "1-3" \\h \\z \\u '
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    placeholder=OxmlElement('w:t'); placeholder.text='打开 Word 后目录将自动更新；如未更新，请右键目录并选择“更新域”。'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    rr=toc.add_run()._r
    rr.append(begin); rr.append(instr); rr.append(sep); rr.append(placeholder); rr.append(end)
    doc.add_page_break()



FN_MARKER_RE = re.compile(r'\[\[FN:([^:\]\n]+):([^\]\n]+)\]\]')


def _note_key(scope, note_id):
    scope=re.sub(r'\s+','-',str(scope or '').strip().lower())
    note_id=str(note_id or '').strip()
    if not scope or not note_id:
        return ''
    return f'{scope}:{note_id}'


def _ensure_footnote_pr(doc):
    """Place true footnotes at the bottom of the page with stable continuous numbering."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for sec in doc.sections:
        sect=sec._sectPr
        old=sect.find(qn('w:footnotePr'))
        if old is not None:
            sect.remove(old)
        fp=OxmlElement('w:footnotePr')
        pos=OxmlElement('w:pos'); pos.set(qn('w:val'),'pageBottom'); fp.append(pos)
        fmt=OxmlElement('w:numFmt'); fmt.set(qn('w:val'),'decimal'); fp.append(fmt)
        start=OxmlElement('w:numStart'); start.set(qn('w:val'),'1'); fp.append(start)
        sect.append(fp)


def _patch_true_footnotes(docx_bytes, note_bank):
    """Convert [[FN:scope:id]] markers to true OOXML footnotes at page bottom.

    note_bank maps `scope:id` -> note text. Unmatched markers are retained visibly
    as superscript `〔注id〕` so the exporter never silently loses a citation.
    """
    import zipfile, copy
    from lxml import etree

    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    PR='http://schemas.openxmlformats.org/package/2006/relationships'
    CT='http://schemas.openxmlformats.org/package/2006/content-types'
    NS={'w':W}
    q=lambda ns,tag:f'{{{ns}}}{tag}'

    src=io.BytesIO(docx_bytes)
    with zipfile.ZipFile(src,'r') as zin:
        files={i.filename:zin.read(i.filename) for i in zin.infolist()}

    doc_root=etree.fromstring(files['word/document.xml'])
    notes_name='word/footnotes.xml'
    if notes_name in files:
        notes_root=etree.fromstring(files[notes_name])
    else:
        notes_root=etree.Element(q(W,'footnotes'),nsmap={'w':W,'r':R})
        for nid,tag in [('-1','separator'),('0','continuationSeparator')]:
            fn=etree.SubElement(notes_root,q(W,'footnote')); fn.set(q(W,'id'),nid); fn.set(q(W,'type'),tag)
            pp=etree.SubElement(fn,q(W,'p')); rr=etree.SubElement(pp,q(W,'r')); etree.SubElement(rr,q(W,tag))

    existing=[]
    for fn in notes_root.findall(q(W,'footnote')):
        try:
            v=int(fn.get(q(W,'id')))
            if v>=1: existing.append(v)
        except Exception: pass
    next_id=max(existing,default=0)+1
    key_to_id={}
    used_keys=[]
    unresolved=[]

    def add_note(key,text):
        nonlocal next_id
        if key in key_to_id:
            return key_to_id[key]
        nid=next_id; next_id+=1; key_to_id[key]=nid; used_keys.append(key)
        fn=etree.SubElement(notes_root,q(W,'footnote')); fn.set(q(W,'id'),str(nid))
        pp=etree.SubElement(fn,q(W,'p'))
        ppr=etree.SubElement(pp,q(W,'pPr'))
        pst=etree.SubElement(ppr,q(W,'pStyle')); pst.set(q(W,'val'),'FootnoteText')
        spacing=etree.SubElement(ppr,q(W,'spacing')); spacing.set(q(W,'after'),'0'); spacing.set(q(W,'line'),'240'); spacing.set(q(W,'lineRule'),'auto')
        rr=etree.SubElement(pp,q(W,'r'))
        rpr=etree.SubElement(rr,q(W,'rPr'))
        rst=etree.SubElement(rpr,q(W,'rStyle')); rst.set(q(W,'val'),'FootnoteReference')
        etree.SubElement(rr,q(W,'footnoteRef'))
        rs=etree.SubElement(pp,q(W,'r')); ts=etree.SubElement(rs,q(W,'t')); ts.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); ts.text=' '
        rr2=etree.SubElement(pp,q(W,'r'))
        rpr2=etree.SubElement(rr2,q(W,'rPr'))
        rf=etree.SubElement(rpr2,q(W,'rFonts')); rf.set(q(W,'ascii'),'Times New Roman'); rf.set(q(W,'hAnsi'),'Times New Roman'); rf.set(q(W,'eastAsia'),'宋体')
        sz=etree.SubElement(rpr2,q(W,'sz')); sz.set(q(W,'val'),'18')
        t=etree.SubElement(rr2,q(W,'t')); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=' '+text
        return nid

    # Snapshot text nodes since we mutate their parents.
    for t in list(doc_root.xpath('.//w:t',namespaces=NS)):
        txt=t.text or ''
        matches=list(FN_MARKER_RE.finditer(txt))
        if not matches: continue
        run=t.getparent()
        if run is None or run.tag!=q(W,'r'): continue
        parent=run.getparent(); idx=parent.index(run)
        original_rpr=run.find(q(W,'rPr'))
        pieces=[]; pos=0
        for m in matches:
            if m.start()>pos: pieces.append(('text',txt[pos:m.start()],None))
            scope,nid=m.group(1).strip(),m.group(2).strip(); key=_note_key(scope,nid)
            if key and key in note_bank:
                pieces.append(('fn',key,add_note(key,note_bank[key])))
            else:
                unresolved.append(key or f'{scope}:{nid}')
                pieces.append(('missing',nid,None))
            pos=m.end()
        if pos<len(txt): pieces.append(('text',txt[pos:],None))
        parent.remove(run)
        offset=0
        for kind,val,note_id in pieces:
            nr=etree.Element(q(W,'r'))
            if original_rpr is not None: nr.append(copy.deepcopy(original_rpr))
            if kind=='text':
                tt=etree.SubElement(nr,q(W,'t'))
                if val.startswith(' ') or val.endswith(' '): tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                tt.text=val
            elif kind=='fn':
                # Reference style is mostly cosmetic; footnoteReference drives Word linkage.
                rpr=nr.find(q(W,'rPr'))
                if rpr is None: rpr=etree.SubElement(nr,q(W,'rPr'))
                rst=etree.SubElement(rpr,q(W,'rStyle')); rst.set(q(W,'val'),'FootnoteReference')
                ref=etree.SubElement(nr,q(W,'footnoteReference')); ref.set(q(W,'id'),str(note_id))
            else:
                rpr=nr.find(q(W,'rPr'))
                if rpr is None: rpr=etree.SubElement(nr,q(W,'rPr'))
                va=etree.SubElement(rpr,q(W,'vertAlign')); va.set(q(W,'val'),'superscript')
                tt=etree.SubElement(nr,q(W,'t')); tt.text=f'〔注{val}〕'
            parent.insert(idx+offset,nr); offset+=1

    # Relationship.
    rel_name='word/_rels/document.xml.rels'
    rel_root=etree.fromstring(files[rel_name])
    rel_type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'
    if not any(x.get('Type')==rel_type for x in rel_root.findall(q(PR,'Relationship'))):
        nums=[]
        for x in rel_root.findall(q(PR,'Relationship')):
            m=re.match(r'rId(\d+)$',x.get('Id',''))
            if m: nums.append(int(m.group(1)))
        rel=etree.SubElement(rel_root,q(PR,'Relationship')); rel.set('Id',f'rId{max(nums,default=0)+1}'); rel.set('Type',rel_type); rel.set('Target','footnotes.xml')

    ct_name='[Content_Types].xml'; ct_root=etree.fromstring(files[ct_name])
    part='/word/footnotes.xml'; ctype='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'
    if not any(x.get('PartName')==part for x in ct_root.findall(q(CT,'Override'))):
        ov=etree.SubElement(ct_root,q(CT,'Override')); ov.set('PartName',part); ov.set('ContentType',ctype)

    # LibreOffice and Word both expect the built-in footnote styles to exist.
    styles_name='word/styles.xml'
    styles_root=etree.fromstring(files[styles_name])
    existing_styles={x.get(q(W,'styleId')) for x in styles_root.findall(q(W,'style'))}
    if 'FootnoteText' not in existing_styles:
        st=etree.SubElement(styles_root,q(W,'style')); st.set(q(W,'type'),'paragraph'); st.set(q(W,'styleId'),'FootnoteText')
        nm=etree.SubElement(st,q(W,'name')); nm.set(q(W,'val'),'Footnote Text')
        bo=etree.SubElement(st,q(W,'basedOn')); bo.set(q(W,'val'),'Normal')
        nx=etree.SubElement(st,q(W,'next')); nx.set(q(W,'val'),'FootnoteText')
        etree.SubElement(st,q(W,'unhideWhenUsed'))
    if 'FootnoteReference' not in existing_styles:
        st=etree.SubElement(styles_root,q(W,'style')); st.set(q(W,'type'),'character'); st.set(q(W,'styleId'),'FootnoteReference')
        nm=etree.SubElement(st,q(W,'name')); nm.set(q(W,'val'),'Footnote Reference')
        bo=etree.SubElement(st,q(W,'basedOn')); bo.set(q(W,'val'),'DefaultParagraphFont')
        rp=etree.SubElement(st,q(W,'rPr')); va=etree.SubElement(rp,q(W,'vertAlign')); va.set(q(W,'val'),'superscript')
        etree.SubElement(st,q(W,'unhideWhenUsed'))
    files[styles_name]=etree.tostring(styles_root,xml_declaration=True,encoding='UTF-8',standalone='yes')

    # Register separator footnotes in settings.xml. LibreOffice requires this
    # to render the note text reliably, not only the note numbers.
    settings_name='word/settings.xml'
    settings_root=etree.fromstring(files[settings_name])
    sfp=settings_root.find(q(W,'footnotePr'))
    if sfp is None:
        sfp=etree.SubElement(settings_root,q(W,'footnotePr'))
    existing_sep={x.get(q(W,'id')) for x in sfp.findall(q(W,'footnote'))}
    for sid in ('-1','0'):
        if sid not in existing_sep:
            ff=etree.SubElement(sfp,q(W,'footnote')); ff.set(q(W,'id'),sid)
    files[settings_name]=etree.tostring(settings_root,xml_declaration=True,encoding='UTF-8',standalone='yes')

    files['word/document.xml']=etree.tostring(doc_root,xml_declaration=True,encoding='UTF-8',standalone='yes')
    files[notes_name]=etree.tostring(notes_root,xml_declaration=True,encoding='UTF-8',standalone='yes')
    files[rel_name]=etree.tostring(rel_root,xml_declaration=True,encoding='UTF-8',standalone='yes')
    files[ct_name]=etree.tostring(ct_root,xml_declaration=True,encoding='UTF-8',standalone='yes')

    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
        for name,data in files.items(): zout.writestr(name,data)
    return out.getvalue(), {'used':used_keys,'unresolved':sorted(set(unresolved))}

def build_manuscript_docx(data):
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_SECTION
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError('缺少 python-docx。请重新运行“启动翻译工具.command”。') from e

    manuscript=data.get('manuscript') or {}
    blocks=manuscript.get('blocks') or []
    if not blocks:
        # Safety fallback: convert existing translated pages to body/heading blocks.
        for b in reflow_blocks(data.get('pages') or [],'target',target=True):
            blocks.append({'type':'chapter' if b['kind']=='heading' else 'body','text':b['text'],'page':b.get('page'),'level':1 if b['kind']=='heading' else 0})

    cleaned=[]
    for b in blocks:
        typ=(b.get('type') or 'body').strip().lower()
        text=re.sub(r'\s+',' ',(b.get('text') or '').strip())
        if typ=='discard' or _hard_noise(text):
            continue
        if text:
            item={'type':typ,'text':text,'page':b.get('page'),'level':b.get('level',0)}
            if typ=='footnote':
                item['note_id']=str(b.get('note_id') or '').strip()
                item['note_scope']=str(b.get('note_scope') or '').strip().lower()
                # Fallback: recover a leading note number if the model omitted note_id.
                if not item['note_id']:
                    m=re.match(r'^\s*(\d{1,4})[.．、)]\s*(.*)$', text)
                    if m:
                        item['note_id']=m.group(1); item['text']=m.group(2).strip()
            cleaned.append(item)
    blocks=cleaned
    if not blocks:
        raise ValueError('书稿重建后没有可导出的正文')

    title=(data.get('title') or data.get('file_name') or '中文译稿').strip()
    title=re.sub(r'\.pdf$','',title,flags=re.I)
    title_blocks=[b for b in blocks if b['type']=='book_title']
    subtitle_blocks=[b for b in blocks if b['type']=='subtitle']
    author_blocks=[b for b in blocks if b['type']=='author']
    note_bank={}
    for b in blocks:
        if b['type']!='footnote': continue
        key=_note_key(b.get('note_scope'), b.get('note_id'))
        if key:
            note_bank[key]=b['text']
    true_footnotes=bool(data.get('true_footnotes', True))
    marker_keys=[]
    for b in blocks:
        if b['type']=='footnote': continue
        for m in FN_MARKER_RE.finditer(b.get('text','')):
            marker_keys.append(_note_key(m.group(1),m.group(2)))
    marker_key_set={k for k in marker_keys if k}
    unmatched_note_keys=[k for k in note_bank if k not in marker_key_set]
    if title_blocks:
        title=title_blocks[0]['text']

    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.7); sec.right_margin=Cm(2.4)
    sec.header_distance=Cm(1.2); sec.footer_distance=Cm(1.2)
    _set_update_fields(doc)
    cp=doc.core_properties; cp.title=title; cp.subject='中文书稿重建版'

    # Base and heading styles.
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.6; normal.paragraph_format.space_after=Pt(0)
    from docx.shared import RGBColor
    for name,size in [('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
        st=doc.styles[name]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)

    # Title page.
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(100); p.paragraph_format.space_after=Pt(20)
    r=p.add_run(title); _set_run_font(r,east_asia='黑体',size_pt=22,bold=True)
    if subtitle_blocks:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22)
        r=p.add_run(subtitle_blocks[0]['text']); _set_run_font(r,east_asia='宋体',size_pt=14)
    for b in author_blocks[:3]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(b['text']); _set_run_font(r,east_asia='楷体',size_pt=12)
    doc.add_page_break()

    # Copyright/front matter collected from the original front pages.
    copyrights=[b for b in blocks if b['type']=='copyright']
    if copyrights:
        for b in copyrights:
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.35
            r=p.add_run(b['text']); _set_run_font(r,east_asia='宋体',size_pt=9)
        doc.add_page_break()

    # A rebuilt dynamic TOC; original toc_entry blocks are intentionally omitted.
    _add_toc_field(doc)

    skip_types={'book_title','subtitle','author','copyright','toc_entry','discard','footnote'}
    last_type=None
    for b in blocks:
        typ=b['type']; text=b['text']
        if typ in skip_types:
            continue
        if typ in ('part','chapter','preface_title','appendix'):
            if last_type is not None:
                if typ in ('chapter','preface_title','appendix'):
                    doc.add_section(WD_SECTION.NEW_PAGE)
                else:
                    doc.add_page_break()
            p=doc.add_paragraph(style='Heading 1')
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(18); p.paragraph_format.keep_with_next=True
            r=p.add_run(text); _set_run_font(r,east_asia='黑体',size_pt=16,bold=True)
        elif typ=='section':
            p=doc.add_paragraph(style='Heading 2')
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before=Pt(13); p.paragraph_format.space_after=Pt(7); p.paragraph_format.keep_with_next=True
            r=p.add_run(text); _set_run_font(r,east_asia='黑体',size_pt=13,bold=True)
        elif typ=='subsection':
            p=doc.add_paragraph(style='Heading 3')
            p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True
            r=p.add_run(text); _set_run_font(r,east_asia='黑体',size_pt=11.5,bold=True)
        elif typ=='body':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Pt(21); p.paragraph_format.line_spacing=1.6; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(text); _set_run_font(r,east_asia='宋体',size_pt=10.5)
        elif typ=='blockquote':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent=Pt(21); p.paragraph_format.right_indent=Pt(21); p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.45
            r=p.add_run(text); _set_run_font(r,east_asia='楷体',size_pt=10)
        elif typ in ('epigraph','dedication'):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.left_indent=Pt(85); p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(16); p.paragraph_format.line_spacing=1.35
            r=p.add_run(text); _set_run_font(r,east_asia='楷体',size_pt=10); r.italic=True
        elif typ in ('figure_caption','table_caption'):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7)
            r=p.add_run(text); _set_run_font(r,east_asia='宋体',size_pt=9.5)
        elif typ=='bibliography':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent=Pt(21); p.paragraph_format.hanging_indent=Pt(21); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(3)
            r=p.add_run(text); _set_run_font(r,east_asia='宋体',size_pt=9.5)
        else:
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Pt(21); p.paragraph_format.line_spacing=1.6
            r=p.add_run(text); _set_run_font(r,east_asia='宋体',size_pt=10.5)
        last_type=typ

    # Never silently lose source notes: notes without a detected body anchor are kept in a review appendix.
    if true_footnotes and unmatched_note_keys:
        doc.add_page_break(); p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('未匹配注释（需人工复核）'); _set_run_font(r,east_asia='黑体',size_pt=16,bold=True)
        for k in unmatched_note_keys:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Pt(21); p.paragraph_format.hanging_indent=Pt(21); p.paragraph_format.line_spacing=1.3
            r=p.add_run(f'{k}  {note_bank[k]}'); _set_run_font(r,east_asia='宋体',size_pt=8.5)

    # Clean output has no inherited source headers/footers; footer contains only page number.
    _ensure_footnote_pr(doc)
    for sec in doc.sections:
        sec.header.is_linked_to_previous=False
        sec.footer.is_linked_to_previous=False
        sec.header.paragraphs[0].text=''
        p=sec.footer.paragraphs[0]; p.text=''; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

    out=io.BytesIO(); doc.save(out); content=out.getvalue()
    if true_footnotes and note_bank:
        content, report=_patch_true_footnotes(content,note_bank)
        return content
    return content




def _chunk_text_as_pages(text, max_chars=7000):
    text=(text or '').strip()
    if not text:
        return []
    paras=[x.strip() for x in re.split(r'\n\s*\n|(?<=。)\s*(?=[一-龥A-Za-z])', text) if x.strip()]
    pages=[]; cur=[]; n=0
    for para in paras:
        if cur and n+len(para)>max_chars:
            pages.append('\n\n'.join(cur)); cur=[]; n=0
        cur.append(para); n+=len(para)+2
    if cur: pages.append('\n\n'.join(cur))
    return pages


def import_docx_bytes(blob):
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError('缺少 python-docx。请重新运行“启动翻译工具.command”。') from e
    doc=Document(io.BytesIO(blob))
    # Prefer a two-column bilingual table generated by this app or similar tools.
    bilingual_rows=[]
    for table in doc.tables:
        if len(table.columns)!=2:
            continue
        for row in table.rows:
            a='\n'.join(p.text.strip() for p in row.cells[0].paragraphs if p.text.strip()).strip()
            b='\n'.join(p.text.strip() for p in row.cells[1].paragraphs if p.text.strip()).strip()
            if not a and not b:
                continue
            if a.lower() in {'original','原文'} and ('译文' in b or 'translation' in b.lower()):
                continue
            if a==b and re.search(r'原\s*PDF\s*第\s*\d+\s*页', a, re.I):
                bilingual_rows.append(('__PAGE__', a))
                continue
            bilingual_rows.append((a,b))
    if bilingual_rows:
        pages=[]; src=[]; tgt=[]; page_n=1
        def flush():
            nonlocal src,tgt,page_n
            if src or tgt:
                pages.append({'n':page_n,'source':'\n\n'.join(src).strip(),'target':'\n\n'.join(tgt).strip()})
                page_n+=1; src=[]; tgt=[]
        for a,b in bilingual_rows:
            if a=='__PAGE__':
                flush()
                m=re.search(r'(\d+)',b); page_n=int(m.group(1)) if m else page_n
                continue
            src.append(a); tgt.append(b)
            if sum(map(len,src))+sum(map(len,tgt))>12000:
                flush()
        flush()
        if any(p['target'] for p in pages):
            return {'kind':'bilingual','pages':pages}
    # Otherwise import the readable document body as an already translated manuscript.
    body=[]
    for p in doc.paragraphs:
        t=p.text.strip()
        if not t:
            body.append('')
            continue
        # Drop the app's own mechanical title/marker lines; AI reconstruction will rebuild hierarchy.
        if re.fullmatch(r'原\s*PDF\s*第\s*\d+\s*页',t,re.I):
            continue
        body.append(t)
    text='\n\n'.join(x for x in body if x!='').strip()
    if not text:
        raise ValueError('Word 中没有读取到可导入的正文或双语对照表。')
    chunks=_chunk_text_as_pages(text)
    return {'kind':'translated','pages':[{'n':i+1,'source':'','target':t} for i,t in enumerate(chunks)]}
class Handler(BaseHTTPRequestHandler):
    def cors(self):
        self.send_header('Access-Control-Allow-Origin','*'); self.send_header('Access-Control-Allow-Headers','Content-Type'); self.send_header('Access-Control-Allow-Methods','GET,POST,OPTIONS')
    def do_OPTIONS(self): self.send_response(204); self.cors(); self.end_headers()
    def json_response(self, obj, status=200):
        out=json.dumps(obj,ensure_ascii=False).encode('utf-8')
        self.send_response(status); self.cors(); self.send_header('Content-Type','application/json; charset=utf-8'); self.end_headers(); self.wfile.write(out)
    def binary_response(self, content, content_type='application/octet-stream', status=200):
        self.send_response(status); self.cors(); self.send_header('Content-Type',content_type); self.send_header('Content-Length',str(len(content))); self.end_headers(); self.wfile.write(content)
    def do_GET(self):
        p=self.path.split('?',1)[0]
        if p=='/api/health':
            self.json_response({'ok':True,'app':APP_ID,'version':APP_VERSION,'port':PORT})
            return
        if p=='/api/key-status':
            provider=normalize_provider(urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query).get('provider',['deepseek'])[0])
            key=load_saved_api_key(provider)
            self.json_response({'saved':bool(key),'provider':provider,'storage':('macOS Keychain' if platform.system()=='Darwin' else ('Windows DPAPI' if platform.system()=='Windows' else 'local secure file')),'masked':('••••'+key[-4:]) if key else ''})
            return
        if p=='/': p='/index.html'
        f=(ROOT/p.lstrip('/')).resolve()
        if not str(f).startswith(str(ROOT)) or not f.is_file(): self.send_error(404); return
        mime='text/html; charset=utf-8' if f.suffix=='.html' else 'text/plain; charset=utf-8'
        self.send_response(200); self.cors(); self.send_header('Content-Type',mime); self.end_headers(); self.wfile.write(f.read_bytes())
    def do_POST(self):
        try:
            n=int(self.headers.get('Content-Length','0')); data=json.loads(self.rfile.read(n) or b'{}')
            if self.path=='/api/process':
                text=call_ai(data); self.json_response({'text':text}); return
            if self.path=='/api/export-docx':
                mode=data.get('mode','translated')
                if mode=='manuscript':
                    content=build_manuscript_docx(data)
                else:
                    content=build_translation_docx(data,bilingual=(mode=='bilingual'))
                self.binary_response(content,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'); return
            if self.path=='/api/import-docx':
                raw=data.get('base64','')
                if ',' in raw: raw=raw.split(',',1)[1]
                try: blob=base64.b64decode(raw, validate=True)
                except Exception: raise ValueError('Word 文件数据无法解析')
                if len(blob)>50*1024*1024: raise ValueError('Word 文件过大（上限 50MB）')
                result=import_docx_bytes(blob)
                self.json_response(result); return
            if self.path=='/api/save-key':
                provider=normalize_provider(data.get('provider')); storage=save_api_key(data.get('api_key'),provider); key=data.get('api_key','').strip(); self.json_response({'saved':True,'storage':storage,'masked':'••••'+key[-4:]}); return
            if self.path=='/api/delete-key':
                provider=normalize_provider(data.get('provider')); delete_saved_api_key(provider); self.json_response({'saved':False,'provider':provider}); return
            self.send_error(404)
        except Exception as e:
            self.json_response({'error':str(e)},500)
    def log_message(self,fmt,*args): print('[server]',fmt%args)

if __name__=='__main__':
    print(f'Scholar PDF Translator v{APP_VERSION}: http://{HOST}:{PORT}')
    print('Ctrl+C 退出。各 AI 平台 API Key 可按平台分别保存在系统安全凭据存储中。')
    ThreadingHTTPServer((HOST,PORT),Handler).serve_forever()