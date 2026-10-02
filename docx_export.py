"""译文分段重排、Word 书稿与双语稿导出、Word 导入。"""

import io, re

from notes import pair_notes, condense_publication_info, SUPERSCRIPT


def _paras(text):
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    # Translation prompts preserve paragraphs with blank lines. If a page contains
    # only line breaks, keep those lines together rather than exploding every line.
    parts = [
        re.sub(r"\s*\n\s*", " ", x).strip()
        for x in re.split(r"\n\s*\n+", text)
        if x.strip()
    ]
    return parts or [text]


def _looks_heading(text):
    t = re.sub(r"\s+", " ", (text or "").strip())
    if not t or len(t) > 120:
        return False
    if re.match(
        r"^(chapter|part|section|introduction|conclusion|preface|foreword|contents|bibliography|references|appendix)\b",
        t,
        re.I,
    ):
        return True
    if re.match(r"^\d+(?:\.\d+){0,3}[\s、.)]+\S+", t):
        return True
    if re.match(r"^[一二三四五六七八九十百]+[、.．]\s*\S+", t):
        return True
    letters = [c for c in t if c.isalpha()]
    if (
        letters
        and len(t) < 80
        and sum(c.isupper() for c in letters) / len(letters) > 0.82
    ):
        return True
    # Common book/article heading shape: short, single-sentence line without end punctuation.
    if (
        len(t) <= 45
        and not re.search(r"[,，。！？.!?;；:：]", t)
        and len(t.split()) <= 9
    ):
        return True
    return False


def _join_across_page(prev, nxt, target=False):
    if not prev or not nxt or _looks_heading(prev) or _looks_heading(nxt):
        return False
    p = prev.rstrip()
    n = nxt.lstrip()
    if target:
        # Chinese prose ending with a sentence terminator normally closes the paragraph.
        # Comma/semicolon/colon or no punctuation at a PDF page boundary is usually continuation.
        return not bool(re.search(r"[。！？!?][”’\"\']?$", p))
    # Strong signal for English/European prose: previous page lacks terminal punctuation,
    # or the next page begins with a lowercase continuation.
    if re.search(r"[.!?][”’\"\']?$", p):
        return False
    if n and n[0].islower():
        return True
    return not bool(re.search(r"[:;][”’\"\']?$", p))


def reflow_blocks(pages, field="target", target=True):
    blocks = []
    for page in pages or []:
        ps = _paras(
            page.get(field) or (page.get("source") if field == "target" else "")
        )
        if not ps:
            continue
        start = len(blocks)
        for x in ps:
            kind = "heading" if _looks_heading(x) else "body"
            blocks.append({"text": x, "kind": kind, "page": page.get("n")})
        # Heal only the boundary created by the original PDF page break.
        if start > 0 and start < len(blocks):
            a = blocks[start - 1]
            b = blocks[start]
            if (
                a["kind"] == "body"
                and b["kind"] == "body"
                and _join_across_page(a["text"], b["text"], target=target)
            ):
                sep = (
                    "" if target and re.search(r"[\u3400-\u9fff]$", a["text"]) else " "
                )
                a["text"] = a["text"].rstrip() + sep + b["text"].lstrip()
                del blocks[start]
    return blocks


def _set_run_font(
    run, latin="Times New Roman", east_asia="宋体", size_pt=None, bold=None
):
    from docx.oxml.ns import qn

    if size_pt is not None:
        from docx.shared import Pt

        run.font.size = Pt(size_pt)
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if bold is not None:
        run.bold = bold
    try:
        from docx.shared import RGBColor

        run.font.color.rgb = RGBColor(0, 0, 0)
    except Exception:
        pass


def _style_doc(doc, bilingual=False):
    from docx.shared import Cm, Pt
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    sec = doc.sections[0]
    if bilingual:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.top_margin = Cm(1.6)
        sec.bottom_margin = Cm(1.6)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
    else:
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.4)
        sec.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(4)
    for name, size in [("Title", 18), ("Heading 1", 15), ("Heading 2", 13)]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
    return sec


def build_translation_docx(data, bilingual=False):
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError(
            "缺少 python-docx。请重新运行“启动翻译工具.command”，新版会自动安装该组件。"
        ) from e

    pages = data.get("pages") or []
    if not pages:
        raise ValueError("没有可导出的页面")
    title = (data.get("title") or data.get("file_name") or "翻译稿").strip()
    title = re.sub(r"\.pdf$", "", title, flags=re.I)
    keep_page_markers = bool(data.get("keep_page_markers", False))
    doc = Document()
    _style_doc(doc, bilingual=bilingual)
    cp = doc.core_properties
    cp.title = title
    cp.subject = "Scholar PDF Translator 导出稿"

    tp = doc.add_paragraph()
    tp.style = "Title"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(title)
    _set_run_font(tr, east_asia="黑体", size_pt=18, bold=True)

    if not bilingual:
        blocks = reflow_blocks(pages, "target", target=True)
        last_page = None
        for b in blocks:
            if keep_page_markers and b.get("page") != last_page:
                m = doc.add_paragraph()
                m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = m.add_run(f'原 PDF 第 {b.get("page")} 页')
                _set_run_font(r, east_asia="宋体", size_pt=8)
                r.font.italic = True
                last_page = b.get("page")
            if b["kind"] == "heading":
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(5)
                r = p.add_run(b["text"])
                _set_run_font(r, east_asia="黑体", size_pt=14, bold=True)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Pt(22)
                p.paragraph_format.line_spacing = 1.5
                r = p.add_run(b["text"])
                _set_run_font(r, east_asia="宋体", size_pt=11)
    else:
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = note.add_run("原文 / 译文对照版")
        _set_run_font(rr, east_asia="黑体", size_pt=10, bold=True)
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [Cm(12.6), Cm(12.6)]
        hdr = table.rows[0].cells
        for i, txt in enumerate(("Original", "中文译文")):
            hdr[i].width = widths[i]
            hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt)
            _set_run_font(r, east_asia="黑体", size_pt=10, bold=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EDEDED")
            hdr[i]._tc.get_or_add_tcPr().append(shd)
        for page in pages:
            src = _paras(page.get("source", ""))
            tgt = _paras(page.get("target") or page.get("source", ""))
            if keep_page_markers:
                cells = table.add_row().cells
                merged = cells[0].merge(cells[1])
                merged.text = ""
                p = merged.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(f'原 PDF 第 {page.get("n")} 页')
                _set_run_font(r, east_asia="宋体", size_pt=8)
                r.font.italic = True
            count = max(len(src), len(tgt), 1)
            for i in range(count):
                cells = table.add_row().cells
                for j, text in enumerate(
                    (src[i] if i < len(src) else "", tgt[i] if i < len(tgt) else "")
                ):
                    cells[j].width = widths[j]
                    cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    p = cells[j].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.25
                    r = p.add_run(text)
                    _set_run_font(r, east_asia="宋体", size_pt=9.5)
        # Repeat header row on subsequent pages.
        trPr = table.rows[0]._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)

    # Footer page numbers using a PAGE field.
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addnext(fld)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


NOISE_PATTERNS = [
    r"^\s*\d{1,5}\s*$",
    r"\b(barcode|call\s*number|library\s*copy|digitized\s*by|scanned\s*by|google\s*books|hathitrust|internet\s*archive)\b",
    r"(图书馆.{0,8}(藏|馆藏|借阅|索书)|馆藏章|扫描.{0,6}(制作|来源|编号)|条形码|条码|索书号)",
    r"^[\W_]{3,}$",
]


def _hard_noise(text):
    t = re.sub(r"\s+", " ", (text or "").strip())
    if not t:
        return True
    if len(t) <= 3 and not re.search(r"[\u3400-\u9fffA-Za-z0-9]", t):
        return True
    return any(re.search(p, t, re.I) for p in NOISE_PATTERNS)


def _set_update_fields(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


FN_MARKER_RE = re.compile(r"\[\[FN:([^:\]\n]+):([^\]\n]+)\]\]")


def _note_key(scope, note_id):
    scope = re.sub(r"\s+", "-", str(scope or "").strip().lower())
    note_id = str(note_id or "").strip()
    if not scope or not note_id:
        return ""
    return f"{scope}:{note_id}"


def _ensure_footnote_pr(doc, numbering="continuous"):
    """页下注放在页面底部。numbering：continuous 全书连续；page 每页重排、带圈数字；chapter 每章（节）重排。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fmt = "decimalEnclosedCircleChinese" if numbering == "page" else "decimal"
    restart = {"page": "eachPage", "chapter": "eachSect"}.get(numbering, "continuous")
    for sec in doc.sections:
        fp = _sect_put(sec._sectPr, "footnotePr", {})
        for tag, val in (
            ("pos", "pageBottom"),
            ("numFmt", fmt),
            ("numStart", "1"),
            ("numRestart", restart),
        ):
            el = OxmlElement("w:" + tag)
            el.set(qn("w:val"), val)
            fp.append(el)


def _patch_true_footnotes(docx_bytes, note_bank, ref_baseline=False):
    """Convert [[FN:scope:id]] markers to true OOXML footnotes at page bottom.

    note_bank maps `scope:id` -> note text. Unmatched markers are retained visibly
    as superscript `〔注id〕` so the exporter never silently loses a citation.
    """
    import zipfile, copy
    from lxml import etree

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    PR = "http://schemas.openxmlformats.org/package/2006/relationships"
    CT = "http://schemas.openxmlformats.org/package/2006/content-types"
    NS = {"w": W}
    q = lambda ns, tag: f"{{{ns}}}{tag}"

    src = io.BytesIO(docx_bytes)
    with zipfile.ZipFile(src, "r") as zin:
        files = {i.filename: zin.read(i.filename) for i in zin.infolist()}

    doc_root = etree.fromstring(files["word/document.xml"])
    notes_name = "word/footnotes.xml"
    if notes_name in files:
        notes_root = etree.fromstring(files[notes_name])
    else:
        notes_root = etree.Element(q(W, "footnotes"), nsmap={"w": W, "r": R})
        for nid, tag in [("-1", "separator"), ("0", "continuationSeparator")]:
            fn = etree.SubElement(notes_root, q(W, "footnote"))
            fn.set(q(W, "id"), nid)
            fn.set(q(W, "type"), tag)
            pp = etree.SubElement(fn, q(W, "p"))
            rr = etree.SubElement(pp, q(W, "r"))
            etree.SubElement(rr, q(W, tag))

    existing = []
    for fn in notes_root.findall(q(W, "footnote")):
        try:
            v = int(fn.get(q(W, "id")))
            if v >= 1:
                existing.append(v)
        except Exception:
            pass
    next_id = max(existing, default=0) + 1
    key_to_id = {}
    used_keys = []
    unresolved = []

    def add_note(key, text):
        nonlocal next_id
        if key in key_to_id:
            return key_to_id[key]
        nid = next_id
        next_id += 1
        key_to_id[key] = nid
        used_keys.append(key)
        fn = etree.SubElement(notes_root, q(W, "footnote"))
        fn.set(q(W, "id"), str(nid))
        pp = etree.SubElement(fn, q(W, "p"))
        ppr = etree.SubElement(pp, q(W, "pPr"))
        pst = etree.SubElement(ppr, q(W, "pStyle"))
        pst.set(q(W, "val"), "FootnoteText")
        rr = etree.SubElement(pp, q(W, "r"))
        rpr = etree.SubElement(rr, q(W, "rPr"))
        rst = etree.SubElement(rpr, q(W, "rStyle"))
        rst.set(q(W, "val"), "FootnoteReference")
        if ref_baseline:
            # 带圈数字在注释区不必上标，与正文字号对齐更清楚
            va = etree.SubElement(rpr, q(W, "vertAlign"))
            va.set(q(W, "val"), "baseline")
        etree.SubElement(rr, q(W, "footnoteRef"))
        # 编号与文字之间空一格；样式中的悬挂缩进让换行后的文字缩进对齐
        rr2 = etree.SubElement(pp, q(W, "r"))
        t = etree.SubElement(rr2, q(W, "t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = " " + text
        return nid

    # Snapshot text nodes since we mutate their parents.
    for t in list(doc_root.xpath(".//w:t", namespaces=NS)):
        txt = t.text or ""
        matches = list(FN_MARKER_RE.finditer(txt))
        if not matches:
            continue
        run = t.getparent()
        if run is None or run.tag != q(W, "r"):
            continue
        parent = run.getparent()
        idx = parent.index(run)
        original_rpr = run.find(q(W, "rPr"))
        pieces = []
        pos = 0
        for m in matches:
            if m.start() > pos:
                pieces.append(("text", txt[pos : m.start()], None))
            scope, nid = m.group(1).strip(), m.group(2).strip()
            key = _note_key(scope, nid)
            if key and key in note_bank:
                pieces.append(("fn", key, add_note(key, note_bank[key])))
            else:
                unresolved.append(key or f"{scope}:{nid}")
                pieces.append(("missing", nid, None))
            pos = m.end()
        if pos < len(txt):
            pieces.append(("text", txt[pos:], None))
        parent.remove(run)
        offset = 0
        for kind, val, note_id in pieces:
            nr = etree.Element(q(W, "r"))
            if original_rpr is not None:
                nr.append(copy.deepcopy(original_rpr))
            if kind == "text":
                tt = etree.SubElement(nr, q(W, "t"))
                if val.startswith(" ") or val.endswith(" "):
                    tt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                tt.text = val
            elif kind == "fn":
                # Reference style is mostly cosmetic; footnoteReference drives Word linkage.
                rpr = nr.find(q(W, "rPr"))
                if rpr is None:
                    rpr = etree.SubElement(nr, q(W, "rPr"))
                rst = etree.Element(q(W, "rStyle"))
                rst.set(q(W, "val"), "FootnoteReference")
                rpr.insert(0, rst)
                ref = etree.SubElement(nr, q(W, "footnoteReference"))
                ref.set(q(W, "id"), str(note_id))
            else:
                rpr = nr.find(q(W, "rPr"))
                if rpr is None:
                    rpr = etree.SubElement(nr, q(W, "rPr"))
                va = etree.SubElement(rpr, q(W, "vertAlign"))
                va.set(q(W, "val"), "superscript")
                tt = etree.SubElement(nr, q(W, "t"))
                tt.text = f"〔注{val}〕"
            parent.insert(idx + offset, nr)
            offset += 1

    # Relationship.
    rel_name = "word/_rels/document.xml.rels"
    rel_root = etree.fromstring(files[rel_name])
    rel_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
    )
    if not any(
        x.get("Type") == rel_type for x in rel_root.findall(q(PR, "Relationship"))
    ):
        nums = []
        for x in rel_root.findall(q(PR, "Relationship")):
            m = re.match(r"rId(\d+)$", x.get("Id", ""))
            if m:
                nums.append(int(m.group(1)))
        rel = etree.SubElement(rel_root, q(PR, "Relationship"))
        rel.set("Id", f"rId{max(nums,default=0)+1}")
        rel.set("Type", rel_type)
        rel.set("Target", "footnotes.xml")

    ct_name = "[Content_Types].xml"
    ct_root = etree.fromstring(files[ct_name])
    part = "/word/footnotes.xml"
    ctype = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
    )
    if not any(x.get("PartName") == part for x in ct_root.findall(q(CT, "Override"))):
        ov = etree.SubElement(ct_root, q(CT, "Override"))
        ov.set("PartName", part)
        ov.set("ContentType", ctype)

    # LibreOffice and Word both expect the built-in footnote styles to exist.
    styles_name = "word/styles.xml"
    styles_root = etree.fromstring(files[styles_name])
    existing_styles = {
        x.get(q(W, "styleId")) for x in styles_root.findall(q(W, "style"))
    }
    if "FootnoteText" not in existing_styles:
        st = etree.SubElement(styles_root, q(W, "style"))
        st.set(q(W, "type"), "paragraph")
        st.set(q(W, "styleId"), "FootnoteText")
        nm = etree.SubElement(st, q(W, "name"))
        nm.set(q(W, "val"), "Footnote Text")
        bo = etree.SubElement(st, q(W, "basedOn"))
        bo.set(q(W, "val"), "Normal")
        nx = etree.SubElement(st, q(W, "next"))
        nx.set(q(W, "val"), "FootnoteText")
        etree.SubElement(st, q(W, "unhideWhenUsed"))
    if "FootnoteReference" not in existing_styles:
        st = etree.SubElement(styles_root, q(W, "style"))
        st.set(q(W, "type"), "character")
        st.set(q(W, "styleId"), "FootnoteReference")
        nm = etree.SubElement(st, q(W, "name"))
        nm.set(q(W, "val"), "Footnote Reference")
        bo = etree.SubElement(st, q(W, "basedOn"))
        bo.set(q(W, "val"), "DefaultParagraphFont")
        rp = etree.SubElement(st, q(W, "rPr"))
        va = etree.SubElement(rp, q(W, "vertAlign"))
        va.set(q(W, "val"), "superscript")
        etree.SubElement(st, q(W, "unhideWhenUsed"))
    files[styles_name] = etree.tostring(
        styles_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    # Register separator footnotes in settings.xml. LibreOffice requires this
    # to render the note text reliably, not only the note numbers.
    settings_name = "word/settings.xml"
    settings_root = etree.fromstring(files[settings_name])
    sfp = settings_root.find(q(W, "footnotePr"))
    if sfp is None:
        sfp = etree.SubElement(settings_root, q(W, "footnotePr"))
    existing_sep = {x.get(q(W, "id")) for x in sfp.findall(q(W, "footnote"))}
    for sid in ("-1", "0"):
        if sid not in existing_sep:
            ff = etree.SubElement(sfp, q(W, "footnote"))
            ff.set(q(W, "id"), sid)
    files[settings_name] = etree.tostring(
        settings_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    files["word/document.xml"] = etree.tostring(
        doc_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    files[notes_name] = etree.tostring(
        notes_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    files[rel_name] = etree.tostring(
        rel_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    files[ct_name] = etree.tostring(
        ct_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    return out.getvalue(), {"used": used_keys, "unresolved": sorted(set(unresolved))}


# ---------- 书稿排版模板 ----------
#
# 所有书稿导出共用一套固定版式。格式写在 Word 样式里，打开文档后可在“样式”面板中
# 统一修改（例如右键“书稿正文”→ 修改），全书随之更新。
#
# 字号对照：二号 22pt，三号 16pt，四号 14pt，小四 12pt，五号 10.5pt，小五 9pt。

LATIN_FONT = "Times New Roman"

BOOK_TEMPLATE = {
    # A4 纸，单位 cm
    "page": {
        "width": 21.0,
        "height": 29.7,
        "top": 2.6,
        "bottom": 2.4,
        "left": 2.8,
        "right": 2.6,
        "header": 1.5,
        "footer": 1.4,
    },
    # 目录收录的标题层级
    "toc_levels": 2,
    # 样式表：ea 中文字体，size 字号（pt），line 行距（倍数），before/after 段前段后（pt），
    # first_chars 首行缩进（字符），left/right/hanging 缩进（pt）
    "styles": [
        dict(name="Normal", ea="宋体", size=12, line=1.5, after=0),
        # 第 1 页：书名、副标题、作者、出版信息，整体居中
        dict(
            name="书名",
            ea="黑体",
            size=22,
            bold=True,
            align="center",
            line=1.3,
            before=150,
            after=14,
        ),
        dict(name="书稿副标题", ea="宋体", size=16, align="center", line=1.3, after=0),
        dict(name="作者", ea="楷体", size=14, align="center", line=1.5, before=40),
        dict(
            name="出版信息",
            ea="宋体",
            size=9,
            align="center",
            line=1.5,
            before=56,
            left=42,
            right=42,
        ),
        # 第 2 页：目录
        dict(name="目录标题", ea="黑体", size=16, align="center", line=1.0, after=24),
        dict(
            name="toc 1",
            sid="TOC1",
            ea="黑体",
            size=12,
            line=1.5,
            before=6,
            after=0,
            toc_tab=True,
        ),
        dict(
            name="toc 2",
            sid="TOC2",
            ea="宋体",
            size=12,
            line=1.5,
            left=24,
            after=0,
            toc_tab=True,
        ),
        dict(
            name="toc 3",
            sid="TOC3",
            ea="宋体",
            size=10.5,
            line=1.5,
            left=48,
            after=0,
            toc_tab=True,
        ),
        # 正文标题
        dict(
            name="Heading 1",
            ea="黑体",
            size=16,
            align="center",
            line=1.3,
            before=24,
            after=24,
            keep=True,
            outline=0,
        ),
        dict(
            name="Heading 2",
            ea="黑体",
            size=14,
            align="left",
            line=1.3,
            before=18,
            after=12,
            keep=True,
            outline=1,
        ),
        dict(
            name="Heading 3",
            ea="黑体",
            size=12,
            align="left",
            line=1.3,
            before=12,
            after=6,
            keep=True,
            outline=2,
        ),
        # 正文与其他段落
        dict(
            name="书稿正文",
            ea="宋体",
            size=12,
            align="justify",
            line=1.5,
            first_chars=2,
        ),
        dict(
            name="引文",
            ea="楷体",
            size=10.5,
            align="justify",
            line=1.5,
            left=24,
            right=24,
            before=6,
            after=6,
        ),
        dict(
            name="题记",
            ea="楷体",
            size=10.5,
            align="right",
            line=1.5,
            left=120,
            before=12,
            after=12,
        ),
        dict(
            name="图表标题",
            ea="黑体",
            size=10.5,
            align="center",
            line=1.3,
            before=6,
            after=6,
        ),
        dict(
            name="参考文献",
            ea="宋体",
            size=10.5,
            align="justify",
            line=1.3,
            left=21,
            hanging=21,
            after=3,
        ),
        # 页下注
        dict(
            name="footnote text",
            sid="FootnoteText",
            ea="宋体",
            size=9,
            align="justify",
            line=1.0,
            left=13.5,
            hanging=13.5,
            after=2,
        ),
        # 页眉页脚
        dict(
            name="Header",
            ea="宋体",
            size=9,
            align="center",
            line=1.0,
            border_bottom=True,
        ),
        dict(name="Footer", ea="宋体", size=9, align="center", line=1.0),
    ],
}

_SECT_ORDER = [
    "headerReference",
    "footerReference",
    "footnotePr",
    "endnotePr",
    "type",
    "pgSz",
    "pgMar",
    "paperSrc",
    "pgBorders",
    "lnNumType",
    "pgNumType",
    "cols",
    "formProt",
    "vAlign",
    "noEndnote",
    "titlePg",
    "textDirection",
    "bidi",
    "rtlGutter",
    "docGrid",
    "printerSettings",
    "sectPrChange",
]


def _sect_put(sectPr, tag, attrs=None):
    """按 OOXML 规定的顺序在 sectPr 中放入（或替换）一个子元素；attrs 为 None 时删除。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for old in sectPr.findall(qn("w:" + tag)):
        sectPr.remove(old)
    if attrs is None:
        return None
    el = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    rank = _SECT_ORDER.index(tag)
    for i, child in enumerate(sectPr):
        name = child.tag.split("}")[-1]
        if name in _SECT_ORDER and _SECT_ORDER.index(name) > rank:
            sectPr.insert(i, el)
            return el
    sectPr.append(el)
    return el


def _set_fonts(rPr, east_asia, latin=LATIN_FONT):
    """直接指定中西文字体，并去掉主题字体属性（主题字体会覆盖直接指定的字体）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + a), None)
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + a), latin)
    rf.set(qn("w:eastAsia"), east_asia)


def _apply_book_template(doc, T=BOOK_TEMPLATE):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    pg = T["page"]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(pg["width"]), Cm(pg["height"])
    sec.top_margin, sec.bottom_margin = Cm(pg["top"]), Cm(pg["bottom"])
    sec.left_margin, sec.right_margin = Cm(pg["left"]), Cm(pg["right"])
    sec.header_distance, sec.footer_distance = Cm(pg["header"]), Cm(pg["footer"])
    text_width_pt = (pg["width"] - pg["left"] - pg["right"]) / 2.54 * 72

    # 文档默认字体也改为直接指定，避免主题字体在 WPS、LibreOffice 中被替换
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            _set_fonts(rpr, "宋体")

    align = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    names = {s.name for s in doc.styles}
    for spec in T["styles"]:
        name = spec["name"]
        if name in names:
            st = doc.styles[name]
        else:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles["Normal"]
            st.quick_style = True
            if spec.get("sid"):
                st.element.set(qn("w:styleId"), spec["sid"])
        font = st.font
        font.size = Pt(spec["size"])
        font.bold = bool(spec.get("bold"))
        font.italic = False
        font.color.rgb = RGBColor(0, 0, 0)
        _set_fonts(st.element.get_or_add_rPr(), spec["ea"])
        pf = st.paragraph_format
        if spec.get("align"):
            pf.alignment = align[spec["align"]]
        pf.line_spacing = spec.get("line", 1.5)
        pf.space_before = Pt(spec.get("before", 0))
        pf.space_after = Pt(spec.get("after", 0))
        pf.left_indent = Pt(spec.get("left", 0))
        pf.right_indent = Pt(spec.get("right", 0))
        if spec.get("hanging"):
            pf.first_line_indent = Pt(-spec["hanging"])
        elif spec.get("first_chars"):
            pf.first_line_indent = Pt(spec["first_chars"] * spec["size"])
        else:
            pf.first_line_indent = Pt(0)
        pf.keep_with_next = bool(spec.get("keep"))
        pf.widow_control = True
        ppr = st.element.get_or_add_pPr()
        if spec.get("first_chars"):
            # 按字符缩进：改字号后首行缩进仍是两个字
            ppr.get_or_add_ind().set(
                qn("w:firstLineChars"), str(int(spec["first_chars"] * 100))
            )
        if name == "Normal":
            # 不对齐文档网格，否则中文行距会被网格放大
            snap = OxmlElement("w:snapToGrid")
            snap.set(qn("w:val"), "0")
            ppr.insert(0, snap)
        if "outline" in spec:
            ol = ppr.find(qn("w:outlineLvl"))
            if ol is None:
                ol = OxmlElement("w:outlineLvl")
                ppr.append(ol)
            ol.set(qn("w:val"), str(spec["outline"]))
        if spec.get("toc_tab"):
            pf.tab_stops.add_tab_stop(
                Pt(text_width_pt - 1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )
        if spec.get("border_bottom"):
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (
                ("val", "single"),
                ("sz", "4"),
                ("space", "4"),
                ("color", "000000"),
            ):
                bottom.set(qn("w:" + k), v)
            bdr.append(bottom)
            ppr.append(bdr)
    # 正文注号：上标
    if "footnote reference" not in names:
        ref = doc.styles.add_style("footnote reference", WD_STYLE_TYPE.CHARACTER)
        ref.element.set(qn("w:styleId"), "FootnoteReference")
        ref.font.superscript = True
        _set_fonts(ref.element.get_or_add_rPr(), "宋体")
    _sort_ppr_children(doc)


_PPR_ORDER = [
    "pStyle",
    "keepNext",
    "keepLines",
    "pageBreakBefore",
    "framePr",
    "widowControl",
    "numPr",
    "suppressLineNumbers",
    "pBdr",
    "shd",
    "tabs",
    "suppressAutoHyphens",
    "kinsoku",
    "wordWrap",
    "overflowPunct",
    "topLinePunct",
    "autoSpaceDE",
    "autoSpaceDN",
    "bidi",
    "adjustRightInd",
    "snapToGrid",
    "spacing",
    "ind",
    "contextualSpacing",
    "mirrorIndents",
    "suppressOverlap",
    "jc",
    "textDirection",
    "textAlignment",
    "textboxTightWrap",
    "outlineLvl",
    "divId",
    "cnfStyle",
    "rPr",
    "sectPr",
    "pPrChange",
]


def _sort_ppr_children(doc):
    """样式里手工加入的段落属性按 OOXML 规定顺序排列，Word 对顺序很严格。"""
    from docx.oxml.ns import qn

    rank = {n: i for i, n in enumerate(_PPR_ORDER)}
    for ppr in doc.styles.element.iter(qn("w:pPr")):
        kids = list(ppr)
        kids.sort(key=lambda el: rank.get(el.tag.split("}")[-1], len(rank)))
        for el in kids:
            ppr.remove(el)
        for el in kids:
            ppr.append(el)


def _fld_run(paragraph, kind, instr=None, dirty=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r = paragraph.add_run()._r
    if instr is not None:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r.append(it)
    else:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        if dirty:
            fc.set(qn("w:dirty"), "true")
        r.append(fc)


def _add_toc(doc, entries, levels):
    """目录页：目录域预先填入各级标题，Word 打开时更新页码。"""
    doc.add_paragraph("目录", style="目录标题")
    entries = [(lv, t) for lv, t in entries if lv <= levels]
    first = doc.add_paragraph(style="toc 1")
    _fld_run(first, "begin", dirty=True)
    _fld_run(first, None, instr=f' TOC \\o "1-{levels}" \\h \\z \\u ')
    _fld_run(first, "separate")
    last = first
    if entries:
        for i, (lv, text) in enumerate(entries):
            p = first if i == 0 else doc.add_paragraph(style=f"toc {lv}")
            if i == 0:
                p.style = doc.styles[f"toc {lv}"]
            p.add_run(text)
            last = p
    else:
        first.add_run("（目录将在 Word 中打开时自动生成）")
    _fld_run(last, "end")


def _page_field(paragraph, roman=False):
    _fld_run(paragraph, "begin")
    _fld_run(paragraph, None, instr=" PAGE \\* ROMAN " if roman else " PAGE ")
    _fld_run(paragraph, "separate")
    paragraph.add_run("I" if roman else "1")
    _fld_run(paragraph, "end")


def _setup_book_sections(doc, running_title):
    """第 1 节书名页（无页眉页码），第 2 节目录（罗马数字页码），
    第 3 节起为正文（页眉为书名，阿拉伯数字页码从 1 开始）。"""
    secs = list(doc.sections)
    for i, sec in enumerate(secs):
        sp = sec._sectPr
        _sect_put(sp, "vAlign", None)
        if i == 1:
            _sect_put(sp, "pgNumType", {"fmt": "upperRoman", "start": 1})
        elif i == 2:
            _sect_put(sp, "pgNumType", {"fmt": "decimal", "start": 1})
        elif i > 2:
            _sect_put(sp, "pgNumType", {"fmt": "decimal"})
        else:
            _sect_put(sp, "pgNumType", None)
        sec.different_first_page_header_footer = False
    # 每一节都写出自己的页眉页脚，不依赖“链接到前一节”，WPS 与 LibreOffice 也能正确显示
    for i, sec in enumerate(secs):
        parts = [sec.header, sec.footer]
        for part in parts:
            part.is_linked_to_previous = False
            for p in part.paragraphs:
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
        if i == 0:
            sec.header.paragraphs[0].style = doc.styles["Footer"]
            continue
        if i >= 2:
            sec.header.paragraphs[0].add_run(running_title)
            _page_field(sec.footer.paragraphs[0])
        else:
            sec.header.paragraphs[0].style = doc.styles["Footer"]
            _page_field(sec.footer.paragraphs[0], roman=True)


def _strip_note_marks(text):
    text = FN_MARKER_RE.sub("", text or "")
    return re.sub(r"[⁰¹²³⁴⁵⁶⁷⁸⁹]+", "", text).strip()


def _drop_empty_note_sections(blocks):
    """注释都转成页下注后，原来的“注释”章只剩标题，删去这样的空章节标题。"""
    top = {"part", "chapter", "preface_title", "appendix"}
    out, i = [], 0
    while i < len(blocks):
        b = blocks[i]
        if b["type"] in top:
            j = i + 1
            while j < len(blocks) and blocks[j]["type"] not in top:
                j += 1
            inner = blocks[i + 1 : j]
            if inner and all(x["type"] in ("footnote", "notes_heading") for x in inner):
                out.extend(inner)
                i = j
                continue
        out.append(b)
        i += 1
    return out


def build_manuscript_docx(data):
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_SECTION
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError(
            "缺少 python-docx。请重新运行“启动翻译工具.command”。"
        ) from e

    manuscript = data.get("manuscript") or {}
    blocks = manuscript.get("blocks") or []
    if not blocks:
        # Safety fallback: convert existing translated pages to body/heading blocks.
        for b in reflow_blocks(data.get("pages") or [], "target", target=True):
            blocks.append(
                {
                    "type": "chapter" if b["kind"] == "heading" else "body",
                    "text": b["text"],
                    "page": b.get("page"),
                    "level": 1 if b["kind"] == "heading" else 0,
                }
            )

    cleaned = []
    for b in blocks:
        typ = (b.get("type") or "body").strip().lower()
        text = re.sub(r"\s+", " ", (b.get("text") or "").strip())
        if typ == "discard" or _hard_noise(text):
            continue
        if text:
            item = {
                "type": typ,
                "text": text,
                "page": b.get("page"),
                "level": b.get("level", 0),
            }
            if typ == "footnote":
                item["note_id"] = str(b.get("note_id") or "").strip()
                item["note_scope"] = str(b.get("note_scope") or "").strip().lower()
                item["origin"] = b.get("origin") or ""
            cleaned.append(item)
    blocks = cleaned
    if not blocks:
        raise ValueError("书稿重建后没有可导出的正文")

    title = (data.get("title") or data.get("file_name") or "中文译稿").strip()
    title = re.sub(r"\.pdf$", "", title, flags=re.I)
    title_blocks = [b for b in blocks if b["type"] == "book_title"]
    subtitle_blocks = [b for b in blocks if b["type"] == "subtitle"]
    author_blocks = [b for b in blocks if b["type"] == "author"]
    source_info = condense_publication_info(blocks)
    blocks, note_bank, note_report = pair_notes(blocks)
    true_footnotes = bool(data.get("true_footnotes", True))
    if true_footnotes:
        blocks = _drop_empty_note_sections(blocks)
    if not true_footnotes:
        # 不转页下注时，把配对好的注号显示为上标数字，注释仍按原位置保留
        blocks = [
            dict(
                b,
                text=FN_MARKER_RE.sub(
                    lambda m: m.group(2).translate(SUPERSCRIPT), b["text"]
                ),
            )
            for b in blocks
        ]
    if title_blocks:
        title = title_blocks[0]["text"]

    T = BOOK_TEMPLATE
    doc = Document()
    _apply_book_template(doc, T)
    _set_update_fields(doc)
    cp = doc.core_properties
    cp.title = title
    cp.subject = "中文书稿重建版"
    if author_blocks:
        cp.author = author_blocks[0]["text"][:200]

    # 第 1 页：书名、副标题、作者、出版信息
    doc.add_paragraph(title, style="书名")
    if subtitle_blocks:
        doc.add_paragraph(subtitle_blocks[0]["text"], style="书稿副标题")
    for i, b in enumerate(author_blocks[:3]):
        p = doc.add_paragraph(b["text"], style="作者")
        if i:
            p.paragraph_format.space_before = Pt(4)
    if source_info:
        doc.add_paragraph(f"译自：{source_info}", style="出版信息")

    # 第 2 页起：目录
    leftover = (
        note_report["unmatched_notes"]
        if true_footnotes
        else note_report["unmatched_notes"]
        + [
            {"page": None, "num": k.split(":", 1)[1], "text": t}
            for k, t in note_bank.items()
        ]
    )
    level_of = {
        "part": 1,
        "chapter": 1,
        "preface_title": 1,
        "appendix": 1,
        "section": 2,
        "subsection": 3,
    }
    toc_entries = [
        (level_of[b["type"]], _strip_note_marks(b["text"]))
        for b in blocks
        if b["type"] in level_of
    ]
    if leftover:
        toc_entries.append((1, "注释" if not true_footnotes else "未能配对的注释"))
    doc.add_section(WD_SECTION.NEW_PAGE)
    _add_toc(doc, toc_entries, T["toc_levels"])
    doc.add_section(WD_SECTION.NEW_PAGE)

    skip_types = {
        "book_title",
        "subtitle",
        "author",
        "copyright",
        "source_info",
        "notes_heading",
        "toc_entry",
        "discard",
        "footnote",
    }
    style_of = {
        "section": "Heading 2",
        "subsection": "Heading 3",
        "body": "书稿正文",
        "blockquote": "引文",
        "epigraph": "题记",
        "dedication": "题记",
        "figure_caption": "图表标题",
        "table_caption": "图表标题",
        "bibliography": "参考文献",
    }
    last_type = None
    for b in blocks:
        typ = b["type"]
        text = b["text"]
        if typ in skip_types:
            continue
        if typ in ("part", "chapter", "preface_title", "appendix"):
            if last_type is not None:
                if typ in ("chapter", "preface_title", "appendix"):
                    doc.add_section(WD_SECTION.NEW_PAGE)
                else:
                    doc.add_page_break()
            p = doc.add_paragraph(text, style="Heading 1")
            if typ == "part":
                # 部标题比章标题大一号，并在页面中上部
                p.paragraph_format.space_before = Pt(120)
                p.runs[0].font.size = Pt(18)
        else:
            doc.add_paragraph(text, style=style_of.get(typ, "书稿正文"))
        last_type = typ

    # 注释不会悄悄丢失：没有找到对应注号的注释集中列在书末，供人工核对
    if leftover:
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph(
            "注释" if not true_footnotes else "未能配对的注释", style="Heading 1"
        )
        for n in leftover:
            where = (
                f"（原书第 {n['page']} 页）" if n.get("page") and true_footnotes else ""
            )
            doc.add_paragraph(f"{n['num']}. {n['text']}{where}", style="footnote text")

    numbering = data.get("footnote_numbering") or "continuous"
    _setup_book_sections(doc, title)
    _ensure_footnote_pr(doc, numbering)

    out = io.BytesIO()
    doc.save(out)
    content = out.getvalue()
    if true_footnotes and note_bank:
        content, _ = _patch_true_footnotes(
            content, note_bank, ref_baseline=(numbering == "page")
        )
    return content


def manuscript_report(data):
    """书稿预览用：注释配对统计与原书信息，与导出 Word 时的结果一致。"""
    blocks = [
        b
        for b in (data.get("manuscript") or {}).get("blocks") or []
        if (b.get("type") or "") != "discard"
    ]
    _, _, report = pair_notes(blocks)
    report["source_info"] = condense_publication_info(blocks)
    return report


# 没有分页信息时，每个导入单元大约相当于一页书稿；单元过长会让书稿重建超出模型的单次输出长度
IMPORT_UNIT_CHARS = 1800


def _split_long_paragraph(para, max_chars):
    """超长段落按句末标点切开，句子本身超长时再硬切。"""
    if len(para) <= max_chars:
        return [para]
    parts, cur = [], ""
    for sent in re.split(r"(?<=[。！？!?；;.])\s*", para):
        if not sent:
            continue
        while len(sent) > max_chars:
            if cur:
                parts.append(cur)
                cur = ""
            parts.append(sent[:max_chars])
            sent = sent[max_chars:]
        if cur and len(cur) + len(sent) > max_chars:
            parts.append(cur)
            cur = sent
        else:
            cur += sent
    if cur:
        parts.append(cur)
    return parts


def _chunk_text_as_pages(text, max_chars=IMPORT_UNIT_CHARS):
    """按段落把连续文本分成约一页长的单元，保留段落边界。"""
    text = (text or "").strip()
    if not text:
        return []
    paras = [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()]
    pages, cur, n = [], [], 0
    for para in paras:
        for piece in _split_long_paragraph(para, max_chars):
            if cur and n + len(piece) > max_chars:
                pages.append("\n\n".join(cur))
                cur, n = [], 0
            cur.append(piece)
            n += len(piece) + 2
    if cur:
        pages.append("\n\n".join(cur))
    return pages


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _paragraph_segments(paragraph):
    """把段落按其中的分页位置切开。

    返回 (段前是否分页, 文字片段列表)。片段之间各隔一个分页；分页来源包括手动分页符、
    段前分页属性，以及 Word 保存时记录的排版分页（lastRenderedPageBreak）。
    """
    el = paragraph._p
    before = False
    ppr = el.find(_W + "pPr")
    if ppr is not None:
        pb = ppr.find(_W + "pageBreakBefore")
        if pb is not None and pb.get(_W + "val") in (None, "1", "true", "on"):
            before = True
    segments, cur = [], []
    for node in el.iter():
        tag = node.tag
        if tag == _W + "t" and node.text:
            cur.append(node.text)
        elif tag == _W + "tab":
            cur.append("\t")
        elif (
            tag == _W + "br" and node.get(_W + "type") == "page"
        ) or tag == _W + "lastRenderedPageBreak":
            if "".join(cur).strip() or segments:
                segments.append("".join(cur))
                cur = []
            else:
                before = True
    segments.append("".join(cur))
    return before, segments


def import_docx_bytes(blob):
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError(
            "缺少 python-docx。请重新运行“启动翻译工具.command”。"
        ) from e
    doc = Document(io.BytesIO(blob))
    # Prefer a two-column bilingual table generated by this app or similar tools.
    bilingual_rows = []
    for table in doc.tables:
        if len(table.columns) != 2:
            continue
        for row in table.rows:
            a = "\n".join(
                p.text.strip() for p in row.cells[0].paragraphs if p.text.strip()
            ).strip()
            b = "\n".join(
                p.text.strip() for p in row.cells[1].paragraphs if p.text.strip()
            ).strip()
            if not a and not b:
                continue
            if a.lower() in {"original", "原文"} and (
                "译文" in b or "translation" in b.lower()
            ):
                continue
            if a == b and re.search(r"原\s*PDF\s*第\s*\d+\s*页", a, re.I):
                bilingual_rows.append(("__PAGE__", a))
                continue
            bilingual_rows.append((a, b))
    if bilingual_rows:
        pages = []
        src = []
        tgt = []
        page_n = 1

        def flush():
            nonlocal src, tgt, page_n
            if src or tgt:
                pages.append(
                    {
                        "n": page_n,
                        "source": "\n\n".join(src).strip(),
                        "target": "\n\n".join(tgt).strip(),
                    }
                )
                page_n += 1
                src = []
                tgt = []

        for a, b in bilingual_rows:
            if a == "__PAGE__":
                flush()
                m = re.search(r"(\d+)", b)
                page_n = int(m.group(1)) if m else page_n
                continue
            src.append(a)
            tgt.append(b)
            if sum(map(len, src)) + sum(map(len, tgt)) > IMPORT_UNIT_CHARS * 2:
                flush()
        flush()
        if any(p["target"] for p in pages):
            return {"kind": "bilingual", "pages": pages}
    # 否则把正文作为已完成的译文导入：优先按 Word 中记录的分页切分，没有分页信息时按约一页的长度切分
    page_texts, cur = [], []
    breaks = 0

    def add(text):
        t = text.strip()
        if t and not re.fullmatch(r"原\s*PDF\s*第\s*\d+\s*页", t, re.I):
            cur.append(t)

    for p in doc.paragraphs:
        before, segments = _paragraph_segments(p)
        if before and cur:
            page_texts.append(cur)
            cur = []
            breaks += 1
        add(segments[0])
        for seg in segments[1:]:
            if cur:
                page_texts.append(cur)
                cur = []
                breaks += 1
            add(seg)
    if cur:
        page_texts.append(cur)
    page_texts = [pg for pg in page_texts if pg]
    if not page_texts:
        raise ValueError("Word 中没有读取到可导入的正文或双语对照表。")
    if breaks >= 2 and len(page_texts) >= 2:
        units = []
        for pg in page_texts:
            units.extend(_chunk_text_as_pages("\n\n".join(pg), IMPORT_UNIT_CHARS * 3))
        unit = "page"
    else:
        units = _chunk_text_as_pages("\n\n".join(t for pg in page_texts for t in pg))
        unit = "chunk"
    return {
        "kind": "translated",
        "unit": unit,
        "pages": [{"n": i + 1, "source": "", "target": t} for i, t in enumerate(units)],
    }
